"""End-to-end: real converter files through the REAL pipeline, scored against
each fixture set's expected.json and defects.json.

    python3 tests/test_end_to_end.py
    ARTPROOF_E2E_REQUIRED=1 python3 tests/test_end_to_end.py   # CI: skips become failures
    ARTPROOF_E2E_SETS=whey python3 tests/test_end_to_end.py    # one set only

Fixture sets (each a folder with expected.json, master_rows.csv and optional
defects.json):
  * tests/fixtures/artwork/        — 4 pancake pouches (run 9ffb350f), 7 scenarios
  * tests/fixtures/artwork/whey/   — 30 whey protein bottles (Git LFS), 1 scenario

No mocks of the reader, no pre-extracted values, no synthetic images. What IS
stubbed, deliberately, is the reference data:

  * The master list. Every set's master_rows.csv is combined into one list (as
    production has one), mutated per scenario, and fetched through the exact
    production path (app._get_sheet_gtin_rows with its retry) and parser
    (app._parse_master_csv) — only app._fetch_sheet_rows, the network edge, is
    replaced. A network failure can never turn this suite green, and
    DEFECT_master_list_unavailable proves the 502 path end to end.
  * The approved panel record (ReadyDoc). Built from expected.json when it
    carries hand-verified serving values; otherwise the lookup answers
    no_panel_values — what ReadyDoc returns today for a file-only approval.
    Prior snapshots / panel versions are None so local history can't leak in.

Defects are planted in that reference data, not in the PDFs: all ProDough label
copy (and, on the bottles, the converter's technical block) is outlined.

Spot colour is OFF by default (DeviceCMYK separations can't be compared
honestly). The pancake set defines its separations in Lab, which can, so the
harness switches the check on for that set only — otherwise its two colour
scenarios and the negative control would test nothing. Sets run in separate
phases so the module-level flag never changes under a running pipeline.

Vision (the panel-crop reader) needs ANTHROPIC_API_KEY. Without it, every
assertion that depends on a panel read is reported SKIP individually, and the
rest still run. Identical vision calls for the same PDF are memoised within a
run so every scenario sees the same read and differs only by reference data.
"""
import copy
import csv
import io
import json
import os
import shutil
import sys
import tempfile
import threading
import time
import urllib.error
from concurrent.futures import ThreadPoolExecutor, as_completed

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

FIX = os.path.join(HERE, 'fixtures', 'artwork')
REQUIRED = bool(os.environ.get('ARTPROOF_E2E_REQUIRED'))
BRAND = {'brand_mode': 'prodough', 'proof_type': 'press', 'wind_direction': ''}
CHECK_KEY = {'gtin': 'gtin', 'net_weight': 'netwt', 'spot_color': 'specs',
             'fda_falcpa': 'fda', 'approved_panel': 'panel', 'item_reference': 'item_reference'}
NFP_FIELDS = ('serving_size_g', 'serving_size_cups', 'servings_per_container', 'serving_size_desc')
MASTER_URL = 'https://app.powder-ops.com/api/products/master.csv?token=e2e'
# Sets whose separations are Lab-defined, so the spot-colour check can run.
SPOT_COLOR_SETS = {'pancake'}


def _skip_all(reason):
    print(f'[{"FAIL" if REQUIRED else "skip"}] end-to-end: {reason}')
    return 1 if REQUIRED else 0


def _discover_sets():
    """[(name, dir)] for every fixture folder holding an expected.json."""
    sets = []
    if os.path.exists(os.path.join(FIX, 'expected.json')):
        sets.append(('pancake', FIX))
    for d in sorted(os.listdir(FIX)) if os.path.isdir(FIX) else []:
        p = os.path.join(FIX, d)
        if os.path.isdir(p) and os.path.exists(os.path.join(p, 'expected.json')):
            sets.append((d, p))
    only = {s.strip() for s in os.environ.get('ARTPROOF_E2E_SETS', '').split(',') if s.strip()}
    return [s for s in sets if not only or s[0] in only]


def _is_lfs_pointer(path):
    with open(path, 'rb') as f:
        return f.read(40).startswith(b'version https://git-lfs')


def _set_unavailable(name, d):
    """Why this set's PDFs can't run, or None."""
    for e in json.load(open(os.path.join(d, 'expected.json')))['files']:
        p = os.path.join(d, e['file'])
        if not os.path.exists(p):
            return f'{e["file"]} missing — see {os.path.relpath(d, HERE)}/README.md'
        if _is_lfs_pointer(p):
            return f'{e["file"]} is a Git LFS pointer, not the PDF — run `git lfs pull`'
    return None


def _preflight(sets):
    if not sets:
        return 'no fixture sets (expected.json) under tests/fixtures/artwork/'
    missing = [t for t in ('pdftoppm', 'tesseract') if not shutil.which(t)]
    if missing:
        return f'PDF/OCR toolchain missing: {", ".join(missing)}'
    try:
        import fitz  # noqa: F401
    except ImportError:
        return 'PyMuPDF (fitz) not installed'
    return None


class Scorer:
    def __init__(self, vision):
        self.vision = vision
        self.fails, self.skips, self.passes = [], [], 0
        self.lock = threading.Lock()

    def check(self, name, cond, needs_vision=False):
        with self.lock:
            if needs_vision and not self.vision:
                self.skips.append(name)
                print(f'[skip] {name}  (needs ANTHROPIC_API_KEY)')
                return
            print(f'[{"ok  " if cond else "FAIL"}] {name}')
            if cond:
                self.passes += 1
            else:
                self.fails.append(name)


def _master_csv(sets, mutations_by_sku=None):
    """Every set's master_rows.csv as one master list, with per-SKU column
    overrides applied, as CSV text."""
    rows, fields = [], None
    for _, d in sets:
        with open(os.path.join(d, 'master_rows.csv'), encoding='utf-8-sig') as f:
            reader = csv.DictReader(f)
            fields = fields or reader.fieldnames
            rows.extend(reader)
    for r in rows:
        for col, val in (mutations_by_sku or {}).get(r['SKU'], {}).items():
            assert col in r, f'defects.json mutates unknown master column {col!r}'
            r[col] = val
    out = io.StringIO()
    w = csv.DictWriter(out, fieldnames=fields)
    w.writeheader()
    w.writerows(rows)
    return out.getvalue()


def _fetch_master(A, csv_text=None, fail_with=None):
    """Fetch the master list through app's production path with the network
    replaced by the local CSV (or by a persistent failure). Returns
    (rows, attempts)."""
    attempts = []

    def _stub(url):
        attempts.append(url)
        if fail_with is not None:
            raise fail_with
        return A._parse_master_csv(csv_text)

    orig_fetch, orig_sleep = A._fetch_sheet_rows, A.time.sleep
    A._fetch_sheet_rows, A.time.sleep = _stub, (lambda s: None)
    A._sheet_cache.clear()
    A._sheet_cache.update({'rows': None, 'fetched_at': 0.0, 'url': ''})
    try:
        rows = A._get_sheet_gtin_rows(force=True)
    finally:
        A._fetch_sheet_rows, A.time.sleep = orig_fetch, orig_sleep
    return rows, attempts


def _install_engine_stubs(pe):
    """Approved panel per run (thread-local), no prior history, and memoised
    vision calls keyed by the source PDF."""
    local = threading.local()
    pe._fetch_approved_panel = lambda gtin=None, sku=None: getattr(local, 'approved', (None, 'not_found'))
    pe._fetch_prior_panel_version = lambda gtin=None, sku=None: None
    pe._fetch_prior_snapshot = lambda gtin=None, sku=None: None

    cache, locks, glock = {}, {}, threading.Lock()

    def _memo(fn, kind):
        def wrapped(img_path, *args):
            key = (kind, os.path.basename(img_path).split('_', 1)[-1], repr(args))
            with glock:
                lk = locks.setdefault(key, threading.Lock())
            with lk:
                if key not in cache:
                    cache[key] = fn(img_path, *args)
                return copy.deepcopy(cache[key])
        return wrapped
    pe._claude_vision_ocr = _memo(pe._claude_vision_ocr, 'vision')
    pe._read_nfp_panel = _memo(pe._read_nfp_panel, 'nfp_crop')
    return local


def _approved(expect, mutate=None):
    """(record, reason) for the stubbed ReadyDoc lookup. With hand-verified
    serving values a real approved record; without, what ReadyDoc answers today
    for a file-only approval."""
    if 'serving_size_g' not in expect:
        return None, 'no_panel_values'
    panel = {'serving_size_g': expect['serving_size_g'],
             'servings_per_container': expect['servings_per_container']}
    panel.update(mutate or {})
    return {'version': 1, 'status': 'approved', 'panel': panel, 'front_callouts': {}}, 'ok'


def _graded(res):
    return [(k, i) for k, c in (res.get('checks') or {}).items() if isinstance(c, dict)
            for i in c.get('issues', []) if i.get('severity') in ('critical', 'warning', 'suspect')]


def _reviews(res):
    # Not graded (a review asks a person to look; it is not a finding), but
    # printed: a file that lands on REVIEW must say why in the log.
    return [(k, i) for k, c in (res.get('checks') or {}).items() if isinstance(c, dict)
            for i in c.get('issues', []) if i.get('severity') == 'review']


def _criticals(res):
    return [(k, i) for k, i in _graded(res) if i['severity'] == 'critical']


def _run():
    sets = _discover_sets()
    why = _preflight(sets)
    if why:
        return _skip_all(why)
    # A set whose PDFs aren't on disk (not yet supplied, or un-pulled LFS
    # pointers) is skipped on its own — loudly — and the rest still run.
    dropped = [(n, _set_unavailable(n, d)) for n, d in sets]
    dropped = [(n, w) for n, w in dropped if w]
    for n, w in dropped:
        print(f'[{"FAIL" if REQUIRED else "skip"}] fixture set "{n}" not run: {w}')
    sets = [s for s in sets if s[0] not in {n for n, _ in dropped}]
    if not sets:
        return 1 if REQUIRED else 0
    for k in ('READYDOC_URL', 'READYDOC_TOKEN', 'GTIN_SHEET_URL', 'MASTER_ALLOW_EXTERNAL'):
        os.environ.pop(k, None)
    os.environ['GTIN_SHEET_URL'] = MASTER_URL
    os.environ['READYDOC_URL'] = 'https://app.powder-ops.com'

    import app as A
    import proof_engine as pe
    A.GTIN_SHEET_CFG_PATH = os.path.join(tempfile.mkdtemp(), 'gtin_sheet_config.json')

    vision = bool(pe.ANTHROPIC_AVAILABLE)
    S = Scorer(vision)
    if not vision:
        print('[skip] vision unavailable (ANTHROPIC_API_KEY not set / anthropic not installed) — '
              'panel-read assertions will be reported individually as SKIP')
    print('fixture sets:', ', '.join(f'{n} ({d})' for n, d in sets))

    # expected: key = path relative to FIX -> entry (+ set name)
    expected, defects = {}, []
    for name, d in sets:
        rel = os.path.relpath(d, FIX)
        for e in json.load(open(os.path.join(d, 'expected.json')))['files']:
            key = e['file'] if rel == '.' else f'{rel}/{e["file"]}'
            expected[key] = dict(e, set=name)
        dp = os.path.join(d, 'defects.json')
        if os.path.exists(dp):
            defects += [dict(x, set=name) for x in json.load(open(dp))['defects']]
    by_sku = {e['sku']: (k, e) for k, e in expected.items()}

    local = _install_engine_stubs(pe)
    work = tempfile.mkdtemp(prefix='artproof_e2e_')

    # ── Build every scenario's inputs (reference data only) ────────────────
    n_master = len(A._parse_master_csv(_master_csv(sets)))
    base_rows, base_attempts = _fetch_master(A, _master_csv(sets))
    S.check(f'combined master list ({n_master} rows) loads through the production fetch path',
            len(base_rows) == n_master and len(base_attempts) == 1)

    jobs = []   # (scenario_id, file_key, rows, approved, set)
    for key, e in expected.items():
        jobs.append(('CLEAN', key, base_rows, _approved(e['expect']), e['set']))
    unavailable = None
    for d in defects:
        if d['sku'] == 'ALL':
            unavailable = d
            continue
        key, e = by_sku[d['sku']]
        rows, _ = _fetch_master(A, _master_csv(sets, {d['sku']: d['mutate'].get('master_row', {})}))
        jobs.append((d['id'], key, rows, _approved(e['expect'], d['mutate'].get('approved_panel')),
                     e['set']))

    if unavailable:
        un_rows, un_attempts = _fetch_master(
            A, fail_with=urllib.error.HTTPError(MASTER_URL, 502, 'Bad Gateway', None, None))
        S.check('DEFECT_master_list_unavailable: the fetch retries 3 times, then yields no rows',
                un_rows == [] and len(un_attempts) == A._MASTER_FETCH_ATTEMPTS)
        blocker = A._master_list_blocker(un_rows, used_upload=False)
        S.check('DEFECT_master_list_unavailable: /upload aborts the run instead of emitting verdicts',
                bool(blocker) and 'aborted' in blocker and '502' in blocker)
        for key, e in expected.items():
            if e['set'] == unavailable['set']:
                jobs.append((unavailable['id'], key, un_rows, (None, 'not_found'), e['set']))

    # ── Run the real pipeline ────────────────────────────────────────────────
    # One run peaks around 3.7 GB (a 920x710mm sheet rasterised at 400 DPI is
    # 162 MP), so default to 2 concurrent runs; ARTPROOF_E2E_WORKERS overrides.
    workers = max(1, int(os.environ.get('ARTPROOF_E2E_WORKERS', '2')))

    def _one(job):
        sid, key, rows, approved, _ = job
        local.approved = approved
        t = time.time()
        res = pe._proof_single(os.path.join(FIX, key), rows, work, brand_config=BRAND,
                               spec_rows=rows)
        return sid, key, res, time.time() - t

    # Each run is printed the moment it finishes, so a timeout or a crash
    # part-way through still leaves every completed result.
    t0 = time.time()
    results = []
    for spot in (True, False):
        phase = [j for j in jobs if (j[4] in SPOT_COLOR_SETS) == spot]
        if not phase:
            continue
        pe.CHECK_SPOT_COLOR = spot
        print(f'\n-- phase: spot colour {"ON" if spot else "OFF (default)"} — {len(phase)} runs', flush=True)
        with ThreadPoolExecutor(max_workers=workers) as ex:
            for fut in as_completed([ex.submit(_one, j) for j in phase]):
                sid, key, res, secs = fut.result()
                results.append((sid, key, res, secs))
                print(f'  [{len(results)}/{len(jobs)}] {sid:<34} {key:<50} '
                      f'{res.get("severity", "?").upper():<10} crit={res.get("critical_count")} '
                      f'warn={res.get("warning_count")} skipped={res.get("checks_skipped")} {secs:.0f}s',
                      flush=True)
                for k, i in _graded(res) + _reviews(res):
                    print(f'      {i["severity"].upper():<8} [{k}] {i["message"][:150]}', flush=True)
                if res.get('vision_allergens'):
                    print(f'      vision allergens: {res["vision_allergens"]}', flush=True)
    print(f'\n{len(results)} pipeline runs in {time.time() - t0:.0f}s\n')

    # A key the API rejects leaves ANTHROPIC_AVAILABLE true while every vision
    # call errors and the pipeline falls back to OCR (run 37844252128, attempt
    # 1). Those numbers look like a vision run but are not one — fail loudly.
    if vision:
        _attempted = [r for _, _, r, _ in results if str(r.get('ocr_preview', '')).startswith('[vision: ')
                      and not str(r.get('ocr_preview', '')).startswith('[vision: not needed')]
        _errored = [r for r in _attempted if str(r.get('ocr_preview', '')).startswith('[vision: error')]
        if _errored:
            print(f'[vision] {len(_errored)} of {len(_attempted)} vision reads errored; first: '
                  + str(_errored[0].get('ocr_preview', ''))[:200])
        S.check(f'vision reads succeed (errored on {len(_errored)} of {len(_attempted)} runs)',
                len(_errored) * 2 < max(1, len(_attempted)))

    # ── Score: clean files ──────────────────────────────────────────────────
    clean_fp = {}
    clean_results = []
    for sid, key, res, _ in results:
        if sid != 'CLEAN':
            continue
        clean_results.append(res)
        exp = expected[key]['expect']
        vals, prov = res.get('panel_values') or {}, res.get('panel_provenance') or {}
        graded = _graded(res)
        clean_fp[expected[key]['set']] = clean_fp.get(expected[key]['set'], 0) + len(graded)
        S.check(f'{key}: zero CRITICAL', not _criticals(res))
        S.check(f'{key}: zero WARNING / SUSPECT findings (got {[(k, i["message"][:50]) for k, i in graded]})',
                not graded)
        if 'item_reference' in exp:
            ir = (res.get('checks') or {}).get('item_reference') or {}
            S.check(f'{key}: item reference reads {exp["item_reference"]} and matches the barcode SKU '
                    f'(got {ir.get("status")} {ir.get("item_ref")})',
                    ir.get('status') == 'OK' and ir.get('item_ref') == exp['item_reference'])
        if 'serving_size_g' in exp:
            S.check(f'{key}: serving size {exp["serving_size_g"]:g}g, {exp.get("serving_size_cups", "-")} cup, '
                    f'{exp["servings_per_container"]:g} servings (got {vals.get("serving_size_g")}, '
                    f'{vals.get("serving_size_cups")}, {vals.get("servings_per_container")})',
                    vals.get('serving_size_g') == exp['serving_size_g']
                    and vals.get('serving_size_cups') == exp.get('serving_size_cups', vals.get('serving_size_cups'))
                    and vals.get('servings_per_container') == exp['servings_per_container'],
                    needs_vision=True)
            S.check(f'{key}: every nutrition-panel field sourced from nfp_crop '
                    f'(got {[prov.get(k) for k in NFP_FIELDS]})',
                    all(prov.get(k) == 'nfp_crop' for k in NFP_FIELDS), needs_vision=True)
        if 'serving_size_desc' in exp:
            want_desc = exp['serving_size_desc'].split('(')[0].strip().lower()
            S.check(f'{key}: serving size description reads "{want_desc}" (got {vals.get("serving_size_desc")!r})',
                    str(vals.get('serving_size_desc') or '').lower().startswith(want_desc),
                    needs_vision=True)
        if 'declared_net_weight_g' in exp:
            S.check(f'{key}: declared net weight {exp["declared_net_weight_g"]:g}g',
                    vals.get('declared_net_weight_g') == exp['declared_net_weight_g'])
        S.check(f'{key}: trim listed as disabled by configuration, not skipped',
                'trim_dimension' in (res.get('checks_disabled') or [])
                and 'trim_dimension' not in (res.get('checks_skipped') or []))
        if exp.get('overall_status'):
            S.check(f'{key}: {exp["overall_status"]}, with every check run (got {res.get("severity")}, '
                    f'skipped {res.get("checks_skipped")})',
                    res.get('severity') == exp['overall_status'].lower() and not res.get('checks_skipped'),
                    needs_vision=True)

    # ── Score: defects ──────────────────────────────────────────────────────
    planted = caught = 0
    negative_ok = None
    for d in defects:
        runs = [r for r in results if r[0] == d['id']]
        exp = d['expect']
        if d['id'] == (unavailable or {}).get('id'):
            for _, key, res, _ in runs:
                vals = res.get('panel_values') or {}
                want = expected[key]['expect']
                S.check(f'{d["id"]} {key}: not CLEAN (got {res.get("severity")})',
                        res.get('severity') != 'clean')
                S.check(f'{d["id"]} {key}: INCOMPLETE (got {res.get("severity")})',
                        res.get('severity') == exp['overall_status'].lower(), needs_vision=True)
                S.check(f'{d["id"]} {key}: serving size still reads {want["serving_size_g"]:g}g / '
                        f'{want["servings_per_container"]:g} with no master list (got '
                        f'{vals.get("serving_size_g")} / {vals.get("servings_per_container")})',
                        vals.get('serving_size_g') == want['serving_size_g']
                        and vals.get('servings_per_container') == want['servings_per_container'],
                        needs_vision=True)
            continue
        (_, key, res, _), = runs
        crits = _criticals(res)
        vision_dep = exp.get('check') in ('approved_panel', 'fda_falcpa')
        if exp['critical_count'] == 0:
            g = _graded(res)
            negative_ok = not g
            S.check(f'{d["id"]}: NEGATIVE CONTROL — zero findings (got {[(k, i["message"][:60]) for k, i in g]})',
                    not g)
            continue
        planted += 1
        ck = CHECK_KEY[exp['check']]
        hit = (len(crits) == 1 and crits[0][0] == ck
               and exp['message_contains'].lower() in crits[0][1]['message'].lower())
        if hit:
            caught += 1
        S.check(f'{d["id"]}: exactly one CRITICAL, on check={exp["check"]}, containing '
                f'"{exp["message_contains"]}" (got {[(k, i["message"][:70]) for k, i in crits]})',
                hit, needs_vision=vision_dep)

    # ── Reader Provenance sheet in the real workbook ─────────────────────────
    import openpyxl
    job = {'id': 'e2e', 'created': time.strftime('%Y-%m-%dT%H:%M:%S'), 'dismissals': {},
           'summary': pe._build_summary(clean_results), 'results': clean_results}
    wb = openpyxl.load_workbook(io.BytesIO(A._generate_report(job).getvalue()))
    S.check('workbook carries a Reader Provenance sheet', 'Reader Provenance' in wb.sheetnames)
    with_panel = [i for i, r in enumerate(clean_results)
                  if 'serving_size_g' in expected[next(k for k, e in expected.items()
                                                       if os.path.basename(k) == r['filename'])]['expect']]
    if 'Reader Provenance' in wb.sheetnames and with_panel:
        ws = wb['Reader Provenance']
        hdr = {ws.cell(row=4, column=c).value: c for c in range(1, ws.max_column + 1)}
        cells = [str(ws.cell(row=5 + i, column=hdr[f]).value) for i in with_panel for f in NFP_FIELDS if f in hdr]
        S.check(f'Reader Provenance sheet shows nfp_crop for every nutrition-panel field on the '
                f'{len(with_panel)} files with verified panel values',
                len(cells) == len(with_panel) * len(NFP_FIELDS) and all('nfp_crop' in c for c in cells),
                needs_vision=True)

    # ── Score line ───────────────────────────────────────────────────────────
    print()
    total_crit = sum(len(_criticals(r)) for s, _, r, _ in results if s == 'CLEAN')
    print(f'SCORE: defects caught {caught} of {planted} planted; negative control '
          f'{"PASS" if negative_ok else "FAIL" if negative_ok is False else "n/a"}; '
          f'CRITICALs on clean files: {total_crit}; '
          f'findings on clean files by set: {clean_fp}'
          + ('' if vision else '   [vision unavailable — panel-read assertions skipped]'))
    print(f'{S.passes} passed, {len(S.fails)} failed, {len(S.skips)} skipped')
    if S.fails:
        print('FAILURES:', *S.fails, sep='\n  - ')
        return 1
    if (S.skips or dropped) and REQUIRED:
        print('ARTPROOF_E2E_REQUIRED is set and assertions or fixture sets were skipped.')
        return 1
    print('End-to-end run passes' + (' (with skips).' if S.skips else '.'))
    return 0


if __name__ == '__main__':
    sys.exit(_run())
