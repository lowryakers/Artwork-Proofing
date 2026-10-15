"""Reader provenance and the silent downgrade (run 9ffb350f).

The causal chain this pack breaks: master list 502 -> fill_weight None and no
approved panel -> _should_run_vision() False -> the geometric panel reader
never ran -> _parse_panel_from_text scraped 65g / 43g off the wrong column of
the layflat -> every check consumed them -> two files filed CLEAN with nothing
read from their panels.

    python3 tests/test_reader_provenance.py

Exits non-zero on any failure. No test framework, no network, no vision API.
"""
import io
import json
import os
import sys
import tempfile
import urllib.error

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import proof_engine as pe  # noqa: E402
import app as A  # noqa: E402


def _run():
    fails = []

    def check(name, cond):
        print(f'[{"ok  " if cond else "FAIL"}] {name}')
        if not cond:
            fails.append(name)

    # ── 1c. The scoped serving-size regex ───────────────────────────────────
    layflat = ('Nutrition Facts\nServing size 3/4 Cup (83g)\nAbout 5 servings per container\n'
               'PROTEIN OPTIONS 19G as packaged 23G with milk\n65g something in another column 43g')
    tp = pe._parse_panel_from_text(layflat)
    check('scoped regex reads 3/4 Cup (83g) -> 83', tp.get('serving_size_g') == 83.0)
    check('scoped regex reads the cup fraction -> 0.75', tp.get('serving_size_cups') == 0.75)
    check('65 (not inside a parenthesis after a measure word) can never match',
          pe._parse_panel_from_text('serving size 65g\n43g') == {})
    check('a unit-declared serving ("4 Cupcakes (63g)") reads grams but no cup value',
          pe._parse_panel_from_text('serving size 4 cupcakes (63g)') == {'serving_size_g': 63.0})
    check('mixed-number cups parse ("1 1/2 cups (120g)")',
          pe._parse_panel_from_text('serving size 1 1/2 cups (120g)').get('serving_size_cups') == 1.5)

    # ── 1a. Rule 2: every panel-bearing file goes to vision ────────────────
    _orig = pe._ocr_needs_vision
    pe._ocr_needs_vision = lambda *a, **k: False
    try:
        complete = ('nutrition facts calories 130 protein 25 contains: milk\n'
                    'serving size 3/4 cup (88g)\nabout 5 servings per container')
        check('a panel-bearing file goes to vision even when the page text looks complete',
              pe._should_run_vision(complete, '', 'PD_pancake.pdf', expects_panel=True) is True)
        check('a file with no panel and complete OCR skips vision',
              pe._should_run_vision(complete, '', 'shipper.pdf', expects_panel=False) is False)
        check('a file with no panel still escalates when the Contains: line is missing',
              pe._should_run_vision('calories 130 protein 25', '', 'shipper.pdf',
                                    expects_panel=False) is True)
    finally:
        pe._ocr_needs_vision = _orig
    check('ProDough files expect a panel by default', pe._expects_nutrition_panel({'brand_mode': 'prodough'}))
    check('an explicit has_nutrition_panel=False opts out',
          not pe._expects_nutrition_panel({'has_nutrition_panel': False}))
    import inspect
    _src = inspect.getsource(pe._should_run_vision)
    check('_should_run_vision no longer references fill_weight or has_approved_panel',
          'fill_weight' not in inspect.signature(pe._should_run_vision).parameters
          and 'has_approved_panel' not in _src)

    # ── 1b / 2. page_regex never supplies the three panel fields; provenance ──
    merged = pe._merge_panel(tp, {})
    check('with no vision read, serving fields are ABSENT (not page_regex values)',
          merged.get('serving_size_g') is None and merged.get('servings_per_container') is None)
    check('the page-regex reads are quarantined under _page_regex for the provenance sheet',
          merged['_page_regex'] == {'serving_size_g': 83.0, 'serving_size_cups': 0.75,
                                    'servings_per_container': 5.0})
    merged2 = pe._merge_panel({'net_weight_g': 454.0, 'unit_count': 21.0}, {'serving_size_g': 88.0})
    check('front-panel fields keep page_regex provenance; vision fields get vision_struct',
          merged2['_src'] == {'net_weight_g': 'page_regex', 'unit_count': 'page_regex',
                              'declared_net_weight_g': 'page_regex', 'serving_size_g': 'vision_struct'})

    # ── NOT RUN: checks lacking a crop-sourced value never produce a verdict ──
    r = pe._check_net_weight({'declared_net_weight_g': 454}, fill_weight_g=454)
    check('net weight with no serving fields -> NOT_RUN (not UNVERIFIED, not PASS)',
          r['status'] == 'NOT_RUN')
    check('NOT_RUN note names both missing fields and points at Reader Provenance',
          any('serving size, servings per container' in n and 'Reader Provenance' in n
              for n in r['notes']))
    check('net weight with no serving fields and no fill stays UNVERIFIED (fill is the limit there)',
          pe._check_net_weight({'declared_net_weight_g': 454}, fill_weight_g=None)['status'] == 'UNVERIFIED')
    check('prep block with an empty panel -> NOT_RUN',
          pe._check_prep_block('Instructions\n1. Mix\n2. Cook', {})['status'] == 'NOT_RUN')
    check('prep block with a serving declaration but no classifiable signal -> UNCLASSIFIED '
          '(a completed advisory read, not a skip)',
          pe._check_prep_block('Instructions\n1. Mix\n2. Cook',
                               {'serving_size_desc': '3/4 Cup'})['status'] == 'UNCLASSIFIED'
          and 'UNCLASSIFIED' not in pe._NOT_VERIFIED_STATUSES)
    check('NOT_RUN is a not-verified status for completeness accounting',
          'NOT_RUN' in pe._NOT_VERIFIED_STATUSES)
    import readydoc as rd
    check('NOT_RUN never files to ReadyDoc as pass', 'NOT_RUN' in rd._UNVERIFIED_STATUSES)

    # ── 3a/3b. CLEAN means every check ran; checks_run / checks_skipped ──────
    res = {'checks': {'gtin': {'status': 'OK', 'issues': []},
                      'netwt': {'status': 'NOT_RUN', 'issues': []},
                      'prep': {'status': 'NOT_RUN', 'issues': []}}}
    pe._recount_result(res)
    check('zero findings from skipped checks -> INCOMPLETE, never CLEAN', res['severity'] == 'incomplete')
    check('checks_run / checks_skipped are recorded',
          res['checks_run'] == ['gtin'] and res['checks_skipped'] == ['netwt', 'prep'])
    res2 = {'checks': {'gtin': {'status': 'OK', 'issues': []}, 'netwt': {'status': 'PASS', 'issues': []}}}
    pe._recount_result(res2)
    check('every check ran and passed -> CLEAN', res2['severity'] == 'clean' and res2['checks_skipped'] == [])
    res3 = {'checks': {'gtin': {'status': 'OK', 'issues': [{'severity': 'warning', 'message': 'x'}]},
                       'netwt': {'status': 'NOT_RUN', 'issues': []}}}
    pe._recount_result(res3)
    check('a real WARNING still outranks INCOMPLETE', res3['severity'] == 'warning')
    res4 = {'checks': {'panel': {'status': 'PANEL_MISSING', 'issues': []}, 'gtin': {'status': 'OK', 'issues': []}}}
    pe._recount_result(res4)
    check('PANEL_MISSING (accurate, not a defect) still means the file is INCOMPLETE, not CLEAN',
          res4['severity'] == 'incomplete')
    s = pe._build_summary([dict(res, filename='a.pdf'), dict(res2, filename='b.pdf')])
    check('_build_summary counts incomplete files separately from clean',
          s['severity_counts']['incomplete'] == 1 and s['severity_counts']['clean'] == 1)

    # ── SUSPECT band recalibration: 105–130 g/cup ───────────────────────────
    for ss in (83, 88, 93):
        r = pe._check_net_weight({'serving_size_g': ss, 'serving_size_cups': 0.75,
                                  'servings_per_container': 5, 'declared_net_weight_g': 454},
                                 fill_weight_g=454)
        check(f'a correct {ss}g / 3/4 cup panel ({ss / 0.75:.0f} g/cup) is PASS with no density finding',
              r['status'] == 'PASS' and not r['issues'])
    r = pe._check_net_weight({'serving_size_g': 198, 'serving_size_cups': 0.75, 'serving_size_desc': '3/4 Cup',
                              'servings_per_container': 2.29, 'declared_net_weight_g': 454}, fill_weight_g=454)
    check('a clean 2x density misread (264 g/cup) is still SUSPECT',
          r['status'] == 'SUSPECT' and any('105–130' in i['message'] for i in r['issues']))

    # ── GTIN: no master list = NOT RUN; wrong barcode on the right artwork = CRITICAL ──
    g = pe._check_gtin('', 'PD_pancake_chocolate.pdf', [], scanned_gtins=['850046726149'])
    check('a decoded GTIN with no master rows -> NOT_RUN, not a silent pass', g.get('status') == 'NOT_RUN')
    rows = [{'gtin': '850046726150', 'flavor': 'chocolate protein pancake mix', 'sku': 'PPM-C'}]
    g = pe._check_gtin('', 'PD_pancake_chocolate.pdf', rows, scanned_gtins=['850046726149'])
    pe._check_gtin_vs_spec_row(g, rows[0])
    crit = [i for i in g['issues'] if i['severity'] == 'critical']
    check('barcode differs from the filename-matched row\'s GTIN -> one CRITICAL naming both',
          len(crit) == 1 and 'does not match' in crit[0]['message']
          and '850046726149' in crit[0]['message'] and '850046726150' in crit[0]['message'])
    check('...and it replaces the vague "not found in the master list" warning',
          not any('not found in the uploaded master list' in i['message'] for i in g['issues']))
    g2 = pe._check_gtin('', 'PD_pancake_chocolate.pdf', [dict(rows[0], gtin='850046726149')],
                        scanned_gtins=['850046726149'])
    pe._check_gtin_vs_spec_row(g2, dict(rows[0], gtin='850046726149'))
    check('barcode equals the row GTIN -> nothing', not g2['issues'])

    # ── Material: same laminate structure described two ways is not a mismatch ──
    check('master-list ply names and a converter abbreviation reduce to the same structure',
          pe._substrate_layers('Karess Soft Touch Matte PET / White 48ga METPET / 3.0mil Clear LLDPE')
          == pe._substrate_layers('PET / MET-PET / LLDPE - 4.5 mil') == ['pet', 'metpet', 'lldpe'])
    check('a genuinely different structure differs',
          pe._substrate_layers('PET / LLDPE') != pe._substrate_layers('PET / MET-PET / LLDPE'))
    check('an unreadable ply never guesses a structure', pe._substrate_layers('Soft Touch / ???') == [])

    # ── NFP: reading nothing is NOT RUN ─────────────────────────────────────
    check('_check_nfp with no calorie/protein read -> NOT_RUN',
          pe._check_nfp('', None)['status'] == 'NOT_RUN')
    check('_check_nfp with values read -> OK',
          pe._check_nfp('', {'nfp': {'calories': 230, 'protein_g': 19}})['status'] == 'OK')

    # ── 3b2. Serving size: the crop's number vs its own verbatim serving line ─
    # CI run 38294631493: the crop returned serving_size_g=370 for a panel that
    # prints "1 bottle (37g)" (whey Chocolate Peanut Butter).
    check('grams parsed from the transcribed serving line',
          [pe._grams_in_serving_text(x) for x in ('1 bottle (37g)', '3/4 Cup (98g)', '2 scoops (31.5 g)', '1 bottle')]
          == [37.0, 98.0, 31.5, None])
    calls = []
    agree = pe._settle_serving_g({'serving_size_g': 98, 'serving_size_text': '3/4 Cup (98g)'},
                                 lambda: calls.append(1), [])
    check('number and text agree -> kept, no re-read', agree['serving_size_g'] == 98 and not calls)
    d = []
    fixed = pe._settle_serving_g({'serving_size_g': 370, 'serving_size_text': '1 bottle (37g)'},
                                 lambda: {'serving_size_g': 37, 'serving_size_text': '1 bottle (37g)'}, d)
    check('370 vs "(37g)" -> one re-read; two readings agree on 37', fixed['serving_size_g'] == 37.0
          and 'two readings agree' in d[0])
    d = []
    unread = pe._settle_serving_g({'serving_size_g': 370, 'serving_size_text': '1 bottle (37g)'},
                                  lambda: {'serving_size_g': 360, 'serving_size_text': '1 bottle (36g)'}, d)
    check('no two readings agree -> left unread (NOT RUN downstream), conflict recorded',
          'serving_size_g' not in unread and 'no two agree' in d[0])
    split = pe._settle_serving_g({'serving_size_g': 370, 'serving_size_text': '1 bottle (37g)'},
                                 lambda: {'serving_size_g': 370, 'serving_size_text': '1 bottle (37g)'}, [])
    check('a 2-2 split is not a decision -> left unread', 'serving_size_g' not in split)

    # Through the real crop reader, API mocked: first reply carries the misread.
    import types
    from PIL import Image as _Img
    _tmp = os.path.join(tempfile.mkdtemp(), 'p.png')
    _Img.new('RGB', (1200, 1600), 'white').save(_tmp)
    _full = {'servings_per_container': 1, 'calories': 130, 'protein_g': 26, 'total_carbohydrate_g': 5,
             'added_sugar_g': 0}
    _replies = iter([json.dumps(dict(_full, serving_size_g=370, serving_size_text='1 bottle (37g)')),
                     json.dumps(dict(_full, serving_size_g=37, serving_size_text='1 bottle (37g)'))])

    class _R:
        def __init__(s, t):
            s.content, s.stop_reason = [types.SimpleNamespace(text=t)], 'end_turn'
    _prev = (pe.ANTHROPIC_AVAILABLE, pe._anthropic)
    os.environ.setdefault('ANTHROPIC_API_KEY', 'test')
    pe.ANTHROPIC_AVAILABLE = True
    pe._anthropic = types.SimpleNamespace(Anthropic=lambda api_key=None: types.SimpleNamespace(
        messages=types.SimpleNamespace(create=lambda **k: _R(next(_replies)))))
    try:
        d = []
        got = pe._read_nfp_panel(_tmp, [0.2, 0.2, 0.8, 0.8], diag=d)
        check(f'crop reader: a 370 misread settles to 37 on re-read ({got.get("serving_size_g")})',
              got.get('serving_size_g') == 37.0 and 'serving_size_text' not in got
              and any('two readings agree' in x for x in d))
    finally:
        pe.ANTHROPIC_AVAILABLE, pe._anthropic = _prev

    # ── 3b3. Panel crop: a tightened crop that reads nothing falls back ────
    # CI run 38789924924: the geometric tightening kept a 1643x439 strip of a
    # tall panel; all three orientations read nothing. Re-read the padded box.
    from PIL import ImageDraw as _Dr
    _tmp2 = os.path.join(tempfile.mkdtemp(), 'q.png')
    _im = _Img.new('RGB', (1200, 1600), (150, 100, 60))
    _Dr.Draw(_im).rectangle((400, 400, 800, 1200), fill='white')
    _im.save(_tmp2)
    _replies2 = iter(['{}', '{}', '{}', json.dumps(dict(_full, serving_size_g=37, serving_size_text='1 bottle (37g)'))])
    pe.ANTHROPIC_AVAILABLE = True
    pe._anthropic = types.SimpleNamespace(Anthropic=lambda api_key=None: types.SimpleNamespace(
        messages=types.SimpleNamespace(create=lambda **k: _R(next(_replies2)))))
    try:
        d = []
        got = pe._read_nfp_panel(_tmp2, [0.33, 0.25, 0.67, 0.75], diag=d)
        check(f'panel crop: tightened read empty -> untightened box re-read ({got.get("serving_size_g")})',
              got.get('serving_size_g') == 37.0 and any('re-reading the untightened box' in x for x in d))
    finally:
        pe.ANTHROPIC_AVAILABLE, pe._anthropic = _prev

    # ── 3b4. A ring field the full read left out gets one focused re-read ──
    # CI run 38808063967: the chocolate pancake crop read 22 fields, no protein,
    # so the protein ring could not be compared with a crop value.
    _no_protein = {k: v for k, v in _full.items() if k != 'protein_g'}
    _replies3 = iter([json.dumps(dict(_no_protein, serving_size_g=83, serving_size_text='3/4 Cup (83g)')),
                      json.dumps({'calories': 290, 'protein_g': 19, 'added_sugar_g': None})])
    _prompts = []

    def _create3(**k):
        _prompts.append(k['messages'][0]['content'][-1]['text'])
        return _R(next(_replies3))
    pe.ANTHROPIC_AVAILABLE = True
    pe._anthropic = types.SimpleNamespace(Anthropic=lambda api_key=None: types.SimpleNamespace(
        messages=types.SimpleNamespace(create=_create3)))
    try:
        d = []
        got = pe._read_nfp_panel(_tmp, [0.2, 0.2, 0.8, 0.8], diag=d)
        check(f'crop: protein missing from the full read -> focused re-read fills it, nothing else '
              f'overwritten ({got.get("protein_g")}, calories {got.get("calories")})',
              got.get('protein_g') == 19 and got.get('calories') == 130 and len(_prompts) == 2
              and _prompts[1] is pe._NFP_RING_FIELDS_PROMPT
              and any('ring-field re-read' in x and 'protein_g' in x for x in d))
    finally:
        pe.ANTHROPIC_AVAILABLE, pe._anthropic = _prev

    # ── 3c. Master-list fetch retries transient failures with backoff ───────
    _orig_fetch, _orig_sleep = A._fetch_sheet_rows, A.time.sleep
    calls = []
    sleeps = []
    A.time.sleep = lambda s: sleeps.append(s)
    try:
        def _flaky(url):
            calls.append(url)
            if len(calls) < 3:
                raise urllib.error.HTTPError(url, 502, 'Bad Gateway', None, None)
            return [{'sku': 'X'}]
        A._fetch_sheet_rows = _flaky
        rows = A._fetch_sheet_rows_with_retry('https://x/master.csv')
        check('two 502s then success -> rows returned after 3 attempts',
              rows == [{'sku': 'X'}] and len(calls) == 3)
        check('backoff slept between attempts', sleeps == [1.0, 2.0])

        calls.clear()
        A._fetch_sheet_rows = lambda url: (_ for _ in ()).throw(
            urllib.error.HTTPError(url, 502, 'Bad Gateway', None, None))
        try:
            A._fetch_sheet_rows_with_retry('https://x/master.csv')
            check('three 502s -> raises', False)
        except urllib.error.HTTPError as e:
            check('three 502s -> raises the last error after 3 attempts', e.code == 502)

        calls.clear()
        def _auth_fail(url):
            calls.append(url)
            raise urllib.error.HTTPError(url, 401, 'Unauthorized', None, None)
        A._fetch_sheet_rows = _auth_fail
        try:
            A._fetch_sheet_rows_with_retry('https://x/master.csv')
        except urllib.error.HTTPError:
            pass
        check('a 401 (not transient) is not retried', len(calls) == 1)
    finally:
        A._fetch_sheet_rows, A.time.sleep = _orig_fetch, _orig_sleep

    # ── 3d. An unavailable master list aborts the run ────────────────────────
    _env = {k: os.environ.get(k) for k in ('GTIN_SHEET_URL', 'READYDOC_URL')}
    _cfg_path, _cache = A.GTIN_SHEET_CFG_PATH, dict(A._sheet_cache)
    tmp = tempfile.mkdtemp()
    A.GTIN_SHEET_CFG_PATH = os.path.join(tmp, 'cfg.json')
    try:
        os.environ['GTIN_SHEET_URL'] = 'https://app.powder-ops.com/api/products/master.csv?token=t'
        os.environ['READYDOC_URL'] = 'https://app.powder-ops.com'
        A._sheet_cache['last_error'] = 'HTTP Error 502: Bad Gateway'
        msg = A._master_list_blocker([], used_upload=False)
        check('configured master list + 0 rows -> run aborted with the fetch error named',
              msg is not None and 'aborted' in msg and '502' in msg and 'no per-file verdicts' in msg.lower())
        check('rows present -> not blocked', A._master_list_blocker([{'sku': 'X'}], False) is None)
        check('an uploaded GTIN file -> not blocked (no master list needed)',
              A._master_list_blocker([], used_upload=True) is None)
        os.environ.pop('GTIN_SHEET_URL')
        check('nothing configured -> not blocked here (the NOT CONFIGURED banner handles that)',
              A._master_list_blocker([], False) is None)
    finally:
        for k, v in _env.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
        A.GTIN_SHEET_CFG_PATH = _cfg_path
        A._sheet_cache.clear(); A._sheet_cache.update(_cache)

    # ── 2 / 3b. The workbook: Reader Provenance sheet + Checks run/skipped ───
    import openpyxl
    job = {'id': 'job1', 'created': '2026-10-05T00:00:00', 'dismissals': {}, 'summary': {},
           'results': [{
               'filename': 'PD_pancake_chocolate.pdf', 'severity': 'incomplete',
               'critical_count': 0, 'warning_count': 0, 'info_count': 0,
               'checks': {'gtin': {'status': 'OK', 'issues': [], 'notes': []},
                          'netwt': {'status': 'NOT_RUN', 'issues': [], 'notes': ['NOT RUN — x']}},
               'checks_run': ['gtin'], 'checks_skipped': ['netwt'],
               'panel_values': {'serving_size_g': None, 'serving_size_cups': None,
                                'servings_per_container': None, 'serving_size_desc': None,
                                'declared_net_weight_g': 454.0, 'unit_count': 21.0},
               'panel_provenance': {'serving_size_g': None, 'serving_size_cups': None,
                                    'servings_per_container': None, 'serving_size_desc': None,
                                    'declared_net_weight_g': 'page_regex', 'unit_count': 'page_regex'},
               'panel_page_regex': {'serving_size_g': 65.0},
           }]}
    buf = A._generate_report(job)
    wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
    check('workbook has a Reader Provenance sheet', 'Reader Provenance' in wb.sheetnames)
    ws = wb['Summary']
    hdrs = [ws.cell(row=4, column=c).value for c in range(1, 8)]
    check('Summary has Checks run / Checks skipped columns',
          hdrs[5] == 'Checks run' and hdrs[6] == 'Checks skipped')
    check('Summary row states INCOMPLETE and names the skipped check',
          ws.cell(row=5, column=2).value == 'INCOMPLETE' and 'netwt' in str(ws.cell(row=5, column=7).value))
    wp = wb['Reader Provenance']
    row = [wp.cell(row=5, column=c).value for c in range(1, 9)]
    check('provenance row marks absent serving fields as NOT RUN',
          'NOT RUN' in str(row[1]) and 'NOT RUN' in str(row[3]))
    check('provenance row labels a page_regex field as PAGE_REGEX, not data',
          'PAGE_REGEX' in str(row[5]))
    check('the unused page-regex serving candidate (65) is listed, flagged, and nowhere else',
          'serving_size_g=65.0' in str(row[7]))

    print()
    if fails:
        print('FAILURES:', *fails, sep='\n  - ')
        return 1
    print('All reader-provenance checks pass.')
    return 0


if __name__ == '__main__':
    sys.exit(_run())
