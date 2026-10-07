"""Disable two checks, fix three, add one (runs 73096882 / 7f8ce432).

    python3 tests/test_disable_and_fixes.py

Exits non-zero on any failure. No test framework, no network, no vision API.
"""
import io
import os
import sys
import urllib.error

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import proof_engine as pe  # noqa: E402
import readydoc as rd  # noqa: E402
import app as A  # noqa: E402


class _StubDoc:
    def __init__(self, objects):
        self._objects = objects

    def xref_length(self):
        return len(self._objects) + 1

    def xref_object(self, xref, compressed=False):
        return self._objects[xref - 1]


def _run():
    fails = []

    def check(name, cond):
        print(f'[{"ok  " if cond else "FAIL"}] {name}')
        if not cond:
            fails.append(name)

    # ── §1/§2: both checks off by default, switchable by env ────────────────
    check('CHECK_SPOT_COLOR defaults off', pe.CHECK_SPOT_COLOR is False)
    check('CHECK_TRIM_DIMENSION defaults off', pe.CHECK_TRIM_DIMENSION is False)
    check('_disabled_checks lists both', pe._disabled_checks() == ['spot_color', 'trim_dimension'])
    os.environ['X_FLAG_T'] = 'true'
    check('_env_flag reads "true"', pe._env_flag('X_FLAG_T', False) is True)
    os.environ.pop('X_FLAG_T')
    check('_env_flag falls back to the default when unset', pe._env_flag('X_FLAG_T', False) is False)

    # The exact DeviceCMYK shape from the whey bottles: never a comparable hex.
    cmyk = ('[ /Separation /PANTONE#20171#20C /DeviceCMYK <<\n    /C0 [ 0 0 0 0 ]\n'
            '    /C1 [ 0 .857627 .99548 0 ]\n'
            '    /FunctionType 2 /N 1 /Domain [ 0 1 ] /Range [ 0 1 0 1 0 1 0 1 ] >> ]')
    got = pe._extract_spot_color_hexes(_StubDoc([cmyk]))['PANTONE#20171#20C']
    check('DeviceCMYK separation -> name-match only (no near-white #FFFEFD, no hex at all)',
          got.get('name_only') is True and 'hex' not in got)

    # ── §3: disabled is not skipped ─────────────────────────────────────────
    res = {'checks': {'gtin': {'status': 'OK', 'issues': []}, 'specs': {'issues': []}},
           'checks_disabled': ['spot_color', 'trim_dimension']}
    pe._recount_result(res)
    check('a file whose only gaps are disabled checks reads CLEAN',
          res['severity'] == 'clean' and res['checks_skipped'] == [])
    import openpyxl
    job = {'id': 'j', 'created': '2026-10-07T00:00:00', 'dismissals': {}, 'summary': {},
           'results': [dict(res, filename='x.pdf', critical_count=0, warning_count=0, info_count=0,
                            checks_run=['gtin', 'specs'])]}
    wb = openpyxl.load_workbook(io.BytesIO(A._generate_report(job).getvalue()))
    ws = wb['Summary']
    check('Summary sheet has both Checks skipped and Checks disabled columns',
          ws.cell(row=4, column=7).value == 'Checks skipped'
          and ws.cell(row=4, column=8).value == 'Checks disabled')
    check('Checks disabled names both and says why; status stays CLEAN',
          'spot_color, trim_dimension (by configuration)' == ws.cell(row=5, column=8).value
          and ws.cell(row=5, column=2).value == 'CLEAN')

    # ── §4: each call-out bound to its own caption ──────────────────────────
    apple_pie = '25G PROTEIN PER SERVING\n1G ADDED SUGAR\n120 CALORIES PER SERVING\n100% GRASS-FED WHEY'
    check('Apple Pie front rings: "1G Added Sugar" is not compared to NFP protein',
          pe._check_protein_claims(apple_pie, 25) == [])
    check('a calories ring next to "protein" is not a protein value either',
          pe._check_protein_claims('PROTEIN 120 cal per serving', 25) == [])
    check('a genuine protein mismatch still fires',
          any(i['severity'] == 'critical' for i in pe._check_protein_claims('PROTEIN 24G per serving', 25)))
    check('a matching protein value is silent', pe._check_protein_claims('PROTEIN 25G per serving', 25) == [])

    # ── §5: the three 404s are told apart ───────────────────────────────────
    def _e(body):
        return urllib.error.HTTPError('u', 404, 'Not Found', None, io.BytesIO(body))
    check('404 no_panel_values -> no_panel_values',
          rd._classify_panel_404(_e(b'{"error": "no_panel_values"}')) == 'no_panel_values')
    check('404 product_not_found -> product_not_found',
          rd._classify_panel_404(_e(b'{"error": "product_not_found"}')) == 'product_not_found')
    check('404 no_panel -> no_panel', rd._classify_panel_404(_e(b'{"error": "no_panel"}')) == 'no_panel')
    check('404 with an unhelpful body -> not_found', rd._classify_panel_404(_e(b'Not Found')) == 'not_found')
    r = pe._check_approved_panel({}, {}, None, None, None, fetch_reason='no_panel_values')
    check('no_panel_values -> PANEL_NO_VALUES with the required wording',
          r['status'] == 'PANEL_NO_VALUES' and r['notes'] == [
              'Panel on file but it carries no typed values (file-only approval) — nutrition cannot '
              'be verified against it. Values need loading in ReadyDoc.'])
    check('product_not_found wording',
          pe._check_approved_panel({}, {}, None, None, None, fetch_reason='product_not_found')['notes']
          == ['No product in ReadyDoc for this GTIN/SKU.'])
    check('no_panel wording',
          pe._check_approved_panel({}, {}, None, None, None, fetch_reason='no_panel')['notes'][0]
          .startswith('Product found; no nutrition panel on file.'))
    check('PANEL_NO_VALUES never verifies (engine and ReadyDoc)',
          'PANEL_NO_VALUES' in pe._NOT_VERIFIED_STATUSES and 'PANEL_NO_VALUES' in rd._UNVERIFIED_STATUSES)
    s = pe._build_summary([{'filename': f'f{i}.pdf', 'severity': 'incomplete',
                            'panel_fetch_reason': 'no_panel_values'} for i in range(3)])
    msg = s['panel_feed_status']['message']
    check('run summary says file-only approval, not "no panel records found"',
          'no typed values (file-only approval)' in msg and 'no panel records found' not in msg)

    # ── §6: an unparseable spec value is never a silent pass ────────────────
    for raw in ('HEX F43BF', 'HEX 4E2CID', 'HX FF9015', 'HEX E613B24'):
        check(f'"{raw}" does not parse as a hex colour', pe._hex_channels(raw) is None)
    csv_text = ('SKU,Product name,GTIN,Fill weight (g),Trim length (mm)\n'
                'PP-X,Thing,850000000001,45O,abc\n')
    row = A._parse_master_csv(csv_text)[0]
    check('master parser records unparseable numeric fields instead of dropping them',
          row['_unparseable'] == {'fill_weight_g': '45O', 'trim_length_mm': 'abc'}
          and row['fill_weight_g'] is None)
    checks = {'netwt': {'status': 'UNVERIFIED', 'issues': [],
                        'notes': ['NOT VERIFIED — no fill weight supplied for this SKU.']}}
    pe._apply_unparseable_spec(checks, row)
    check('net weight with an unparseable fill weight -> NOT_RUN + "could not evaluate"',
          checks['netwt']['status'] == 'NOT_RUN'
          and any('Could not evaluate — spec value "45O"' in i['message'] for i in checks['netwt']['issues'])
          and not any('no fill weight supplied' in n for n in checks['netwt']['notes']))

    # ── §7: converter item reference vs the barcode's SKU ───────────────────
    rows = [{'gtin': '850079939387', 'sku': 'WHY-BTL-DCH'}]
    r7 = pe._check_item_reference('Cust. Item Ref: BEF-BTL-DCH PRODOUGH DOUBLE CH…', ['850079939387'], rows)
    check('7 Oct proof (filed as BEF-BTL-DCH, barcode is whey) -> exactly the required CRITICAL',
          [i['message'] for i in r7['issues']] == [
              'Item-reference mismatch — the proof sheet is filed as "BEF-BTL-DCH" but the artwork\'s '
              'barcode (850079939387) is WHY-BTL-DCH. The converter may file or reorder this artwork '
              'against the wrong product.'])
    r5 = pe._check_item_reference('Cust. Item Ref: WHY-BTL-DCH PRODOUGH DOUBLE CH…', ['850079939387'], rows)
    check('5 Oct proof (WHY-BTL-DCH) -> OK, no findings', r5['status'] == 'OK' and not r5['issues'])
    check('no item reference on the sheet -> check does not apply (None)',
          pe._check_item_reference('no reference here', ['850079939387'], rows) is None)
    check('item reference present but barcode unresolved -> NOT_RUN',
          pe._check_item_reference('Cust. Item Ref: BEF-BTL-DCH', ['000000000000'], rows)['status'] == 'NOT_RUN')

    print()
    if fails:
        print('FAILURES:', *fails, sep='\n  - ')
        return 1
    print('All disable-and-fixes checks pass.')
    return 0


if __name__ == '__main__':
    sys.exit(_run())
